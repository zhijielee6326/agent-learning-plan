"""
保险公司产品精算智能体 - FastAPI后端服务
提供条款检索、流式生成、导出、LLM智能对话、条款起草等API接口
"""

from fastapi import FastAPI, Query, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse, Response
from pydantic import BaseModel
from typing import List, Optional
import json
import os
import sys

# 确保可以导入同目录模块
sys.path.insert(0, os.path.dirname(__file__))

from product_store import get_product_store
from clause_generator import ClauseGenerator
from llm_agent import LLMAgent
from vector_store import init_vector_store
from clause_draft_agent import ClauseDraftAgent

app = FastAPI(
    title="保险公司产品精算智能体",
    description="内部产品条款检索、生成与溯源系统（RAG + 向量检索 + LLM）",
    version="3.0.0"
)

# 跨域支持
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["content-disposition"],
)

# 全局实例
store = get_product_store()
generator = ClauseGenerator(store)

# 初始化向量库并导入产品数据
print("[Main] 初始化 Qdrant 向量库...")
vector_store = init_vector_store(store.products)

# 初始化 LLM Agent（注入向量检索）
llm_agent = LLMAgent(store=store, vector_store=vector_store)

# 初始化条款生成Agent（注入向量检索）
clause_agent = ClauseDraftAgent(store=store, vector_store=vector_store)


# ============ 请求模型 ============

class QueryRequest(BaseModel):
    query: str
    mode: str = "standard"  # standard, brief, explain

class ExportRequest(BaseModel):
    query: str
    mode: str = "standard"
    format: str = "markdown"  # markdown, json

class CombinedRequest(BaseModel):
    product_names: List[str]
    mode: str = "standard"

class ChatRequest(BaseModel):
    query: str
    session_id: str = "default"
    mode: str = "standard"

class EditRequest(BaseModel):
    selected_text: str
    instruction: str
    context: str = ""

class ClauseDraftRequest(BaseModel):
    query: str  # 如 "帮我生成一份人身意外伤害保险的产品条款"

    def is_valid(self) -> bool:
        return len(self.query.strip()) >= 4

class WordExportRequest(BaseModel):
    content: str  # Markdown内容
    title: str = "保险产品条款"


# ============ 原有API接口（保留兼容） ============

@app.get("/api/products")
async def list_products():
    """获取所有可用备案产品列表"""
    products = store.list_products()
    return {"products": products, "total": len(products)}


@app.get("/api/products/{product_id}")
async def get_product(product_id: str):
    """获取单个产品详情"""
    product = store.get_product(product_id)
    if not product:
        raise HTTPException(status_code=404, detail="产品未找到")
    return product


@app.post("/api/search")
async def search_clauses(request: QueryRequest):
    """检索匹配条款"""
    result = store.search_sections(request.query)
    return result.to_dict()


@app.post("/api/generate/stream")
async def generate_stream(request: QueryRequest):
    """流式生成条款（SSE）- 原有规则引擎"""
    async def event_stream():
        async for chunk in generator.generate_stream(request.query, request.mode):
            yield chunk

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


@app.post("/api/generate")
async def generate_static(request: QueryRequest):
    """静态生成条款（一次性返回）"""
    result = generator.generate_static(request.query, request.mode)
    return result


@app.post("/api/combined")
async def generate_combined(request: CombinedRequest):
    """组合多产品条款"""
    combined = store.generate_combined_clauses(request.product_names)
    return combined


@app.post("/api/export")
async def export_clauses(request: ExportRequest):
    """导出条款为Markdown或JSON"""
    result = generator.generate_static(request.query, request.mode)

    if request.format == "json":
        return JSONResponse(
            content=result,
            media_type="application/json",
            headers={
                "Content-Disposition": "attachment; filename=clauses.json"
            }
        )
    else:
        # Markdown导出
        md_content = result["markdown"]
        md_content += "\n\n---\n\n## 引用来源清单\n\n"
        for idx, src in enumerate(result["sources"], 1):
            md_content += f"{idx}. {src}\n"

        return JSONResponse(
            content={"markdown": md_content, "sources": result["sources"]},
            headers={
                "Content-Disposition": "attachment; filename=clauses.md"
            }
        )


@app.get("/api/chapters")
async def list_chapters():
    """获取标准章节列表"""
    from product_store import STANDARD_CHAPTER_ORDER
    return {"chapters": STANDARD_CHAPTER_ORDER}


# ============ LLM智能对话接口 ============

@app.post("/api/chat/stream")
async def chat_stream(request: ChatRequest):
    """LLM智能对话 - 流式输出（SSE）- 支持条款生成意图路由"""
    if not request.query.strip():
        raise HTTPException(status_code=400, detail="请输入您的问题")
    # 检测是否为条款生成意图
    if clause_agent.is_clause_generation_intent(request.query):
        async def clause_event_stream():
            async for chunk in clause_agent.draft_stream(request.query):
                yield chunk
        return StreamingResponse(
            clause_event_stream(),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no"
            }
        )

    async def event_stream():
        async for chunk in llm_agent.chat_stream(request.query, request.session_id, request.mode):
            yield chunk

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


@app.post("/api/chat")
async def chat_static(request: ChatRequest):
    """LLM智能对话 - 非流式"""
    try:
        result = await llm_agent.chat_static(request.query, request.session_id)
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/chat/clear")
async def clear_chat(session_id: str = "default"):
    """清除对话历史"""
    llm_agent.clear_conversation(session_id)
    return {"status": "ok", "message": "对话历史已清除"}


# ============ AI辅助编辑接口 ============

@app.post("/api/edit/stream")
async def edit_stream(request: EditRequest):
    """AI辅助编辑 - 流式输出（SSE）"""
    async def event_stream():
        async for chunk in llm_agent.edit_stream(
            request.selected_text, request.instruction, request.context
        ):
            yield chunk

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


# ============ 条款生成专用接口 ============

@app.post("/api/clause-draft/stream")
async def clause_draft_stream(request: ClauseDraftRequest):
    """条款生成工作流 - 流式输出（SSE）"""
    if not request.is_valid():
        raise HTTPException(status_code=400, detail="请求内容过短，请描述您需要生成的保险产品类型")
    async def event_stream():
        async for chunk in clause_agent.draft_stream(request.query):
            yield chunk

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


@app.post("/api/export/word")
async def export_word(request: WordExportRequest):
    """将Markdown内容导出为Word (.docx) 文件"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import re
        from datetime import datetime

        doc = Document()

        # 设置默认字体
        style = doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)

        # 标题
        title_para = doc.add_heading(request.title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 副标题 - 生成日期
        date_para = doc.add_paragraph(f"生成日期：{datetime.now().strftime('%Y年%m月%d日')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 添加分隔线（而非分页，避免空白页）
        doc.add_paragraph("—")

        # 解析 Markdown 内联格式（粗体、斜体、溯源标注）
        def add_inline_text(para, text, base_size=12):
            """解析并添加内联格式文本到段落"""
            # 拆分溯源标注 [...]
            parts = re.split(r'(\[[^\]]+\])', text)
            for part in parts:
                if re.match(r'\[.+\]', part):
                    run = para.add_run(part)
                    run.font.color.rgb = RGBColor(0, 102, 204)
                    run.font.size = Pt(9)
                else:
                    # 拆分 **粗体** 和 *斜体*
                    segments = re.split(r'(\*\*.+?\*\*|\*.+?\*)', part)
                    for seg in segments:
                        if seg.startswith('**') and seg.endswith('**'):
                            run = para.add_run(seg[2:-2])
                            run.font.size = Pt(base_size)
                            run.bold = True
                        elif seg.startswith('*') and seg.endswith('*'):
                            run = para.add_run(seg[1:-1])
                            run.font.size = Pt(base_size)
                            run.italic = True
                        else:
                            if seg:
                                run = para.add_run(seg)
                                run.font.size = Pt(base_size)

        # 解析Markdown内容
        lines = request.content.split('\n')
        in_list = False
        for line in lines:
            line = line.rstrip()

            if not line.strip():
                continue

            # H1 标题
            if line.startswith('# ') and not line.startswith('## '):
                heading_text = line[2:].replace('**', '').strip()
                doc.add_heading(heading_text, level=1)
            # H2 标题
            elif line.startswith('## ') and not line.startswith('### '):
                heading_text = line[3:].replace('**', '').strip()
                doc.add_heading(heading_text, level=2)
            # H3 标题
            elif line.startswith('### ') and not line.startswith('#### '):
                heading_text = line[4:].replace('**', '').strip()
                doc.add_heading(heading_text, level=3)
            # H4 标题
            elif line.startswith('#### '):
                heading_text = line[5:].replace('**', '').strip()
                h = doc.add_heading(heading_text, level=4)

            # 引用块（溯源标注）
            elif line.startswith('> '):
                quote_text = line[2:].strip()
                para = doc.add_paragraph()
                add_inline_text(para, quote_text, base_size=10)
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0, 102, 204)

            # 水平线
            elif line.strip() == '---':
                doc.add_paragraph("—" * 40)

            # 无序列表
            elif line.strip().startswith('- ') or line.strip().startswith('• '):
                text = re.sub(r'^[\s]*[-•]\s+', '', line)
                para = doc.add_paragraph(style='List Bullet')
                add_inline_text(para, text)

            # 有序列表
            elif re.match(r'^\s*\d+[\.、]\s', line):
                text = re.sub(r'^\s*\d+[\.、]\s+', '', line)
                para = doc.add_paragraph(style='List Number')
                add_inline_text(para, text)

            # 普通段落
            else:
                para = doc.add_paragraph()
                add_inline_text(para, line)

        # 保存到内存
        from io import BytesIO
        from urllib.parse import quote
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        raw_filename = f"{request.title}_{timestamp}.docx"
        encoded_filename = quote(raw_filename)

        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
            }
        )
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="python-docx 未安装，请运行：pip install python-docx"
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Word导出失败：{str(e)}")


@app.get("/api/health")
async def health_check():
    """健康检查"""
    return {
        "status": "ok",
        "products_loaded": len(store.products),
        "service": "保险公司产品精算智能体",
        "llm_model": llm_agent.model,
        "vector_db": "Qdrant",
        "embedding_model": os.getenv("EMBEDDING_MODEL", "qwen3-embedding:8b"),
        "version": "3.0.0"
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)